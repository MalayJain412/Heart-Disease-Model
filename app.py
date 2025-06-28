import os
import pickle
import pandas as pd
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import io
import json
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors
import matplotlib.pyplot as plt
import matplotlib
matplotlib.use('Agg')
import seaborn as sns
from collections import Counter
import base64
from config import config
import time

# Create Flask app
app = Flask(__name__, template_folder='app/templates')

# Load configuration
config_name = os.environ.get('FLASK_ENV', 'development')
app.config.from_object(config[config_name])

db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

# Load the pre-trained model
def load_model():
    try:
        model_path = app.config['MODEL_PATH']
        if not os.path.exists(model_path):
            raise FileNotFoundError(f"Model file not found: {model_path}")
        
        with open(model_path, 'rb') as file:
            model = pickle.load(file)
        print(f"Model loaded successfully from {model_path}")
        return model
    except Exception as e:
        print(f"Error loading model: {e}")
        return None

# Initialize model (will be loaded when app starts)
model = None

# Load model immediately after app configuration
print("Loading machine learning model...")
model = load_model()
if model is None:
    print("WARNING: Model could not be loaded. Predictions will not work.")
else:
    print("Model loaded successfully!")

# Database Models
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(120), nullable=False)
    role = db.Column(db.String(20), nullable=False)  # 'doctor', 'dean', 'chairman'
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    patients = db.relationship('Patient', backref='doctor', lazy=True)

class Patient(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    gender = db.Column(db.String(10), nullable=False)
    age = db.Column(db.Integer, nullable=False)
    chest_pain = db.Column(db.Boolean, default=False)
    shortness_of_breath = db.Column(db.Boolean, default=False)
    fatigue = db.Column(db.Boolean, default=False)
    systolic = db.Column(db.Integer, nullable=False)
    diastolic = db.Column(db.Integer, nullable=False)
    heart_rate = db.Column(db.Integer, nullable=False)
    lung_sounds = db.Column(db.Boolean, default=False)
    cholesterol = db.Column(db.Integer, nullable=False)
    ldl = db.Column(db.Integer, nullable=False)
    hdl = db.Column(db.Integer, nullable=False)
    diabetes = db.Column(db.Boolean, default=False)
    atrial_fibrillation = db.Column(db.Boolean, default=False)
    rheumatic_fever = db.Column(db.Boolean, default=False)
    mitral_stenosis = db.Column(db.Boolean, default=False)
    aortic_stenosis = db.Column(db.Boolean, default=False)
    tricuspid_stenosis = db.Column(db.Boolean, default=False)
    pulmonary_stenosis = db.Column(db.Boolean, default=False)
    dilated_cardiomyopathy = db.Column(db.Boolean, default=False)
    hypertrophic_cardiomyopathy = db.Column(db.Boolean, default=False)
    drug_use = db.Column(db.Boolean, default=False)
    fever = db.Column(db.Boolean, default=False)
    chills = db.Column(db.Boolean, default=False)
    alcoholism = db.Column(db.Boolean, default=False)
    hypertension = db.Column(db.Boolean, default=False)
    fainting = db.Column(db.Boolean, default=False)
    dizziness = db.Column(db.Boolean, default=False)
    smoking = db.Column(db.Boolean, default=False)
    obesity = db.Column(db.Boolean, default=False)
    murmur = db.Column(db.Boolean, default=False)
    prediction_result = db.Column(db.String(100), nullable=False)
    doctor_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)

@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))

# Routes
@app.route('/')
def index():
    if current_user.is_authenticated:
        if current_user.role == 'doctor':
            return redirect(url_for('doctor_dashboard'))
        elif current_user.role == 'dean':
            return redirect(url_for('dean_dashboard'))
        elif current_user.role == 'chairman':
            return redirect(url_for('chairman_dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        
        if user and check_password_hash(user.password_hash, password):
            login_user(user)
            flash('Login successful!', 'success')
            return redirect(url_for('index'))
        else:
            flash('Invalid username or password', 'error')
    
    return render_template('auth/login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))

@app.route('/change_password', methods=['GET', 'POST'])
@login_required
def change_password():
    if request.method == 'POST':
        current_password = request.form.get('current_password')
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        
        if not check_password_hash(current_user.password_hash, current_password):
            flash('Current password is incorrect', 'error')
        elif new_password != confirm_password:
            flash('New passwords do not match', 'error')
        else:
            current_user.password_hash = generate_password_hash(new_password)
            db.session.commit()
            flash('Password changed successfully!', 'success')
            return redirect(url_for('index'))
    
    return render_template('auth/change_password.html')

# Doctor Routes
@app.route('/doctor/dashboard')
@login_required
def doctor_dashboard():
    if current_user.role != 'doctor':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.filter_by(doctor_id=current_user.id).order_by(Patient.timestamp.desc()).all()
    return render_template('doctor/dashboard.html', patients=patients)

@app.route('/doctor/add_patient', methods=['GET', 'POST'])
@login_required
def add_patient():
    if current_user.role != 'doctor':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    if model is None:
        flash('Model not loaded. Please contact administrator.', 'error')
        return redirect(url_for('doctor_dashboard'))
    
    if request.method == 'POST':
        # Collect form data
        patient_data = {
            'name': request.form.get('name'),
            'gender': request.form.get('gender'),
            'age': int(request.form.get('age')),
            'chest_pain': bool(request.form.get('chest_pain')),
            'shortness_of_breath': bool(request.form.get('shortness_of_breath')),
            'fatigue': bool(request.form.get('fatigue')),
            'systolic': int(request.form.get('systolic')),
            'diastolic': int(request.form.get('diastolic')),
            'heart_rate': int(request.form.get('heart_rate')),
            'lung_sounds': bool(request.form.get('lung_sounds')),
            'cholesterol': int(request.form.get('cholesterol')),
            'ldl': int(request.form.get('ldl')),
            'hdl': int(request.form.get('hdl')),
            'diabetes': bool(request.form.get('diabetes')),
            'atrial_fibrillation': bool(request.form.get('atrial_fibrillation')),
            'rheumatic_fever': bool(request.form.get('rheumatic_fever')),
            'mitral_stenosis': bool(request.form.get('mitral_stenosis')),
            'aortic_stenosis': bool(request.form.get('aortic_stenosis')),
            'tricuspid_stenosis': bool(request.form.get('tricuspid_stenosis')),
            'pulmonary_stenosis': bool(request.form.get('pulmonary_stenosis')),
            'dilated_cardiomyopathy': bool(request.form.get('dilated_cardiomyopathy')),
            'hypertrophic_cardiomyopathy': bool(request.form.get('hypertrophic_cardiomyopathy')),
            'drug_use': bool(request.form.get('drug_use')),
            'fever': bool(request.form.get('fever')),
            'chills': bool(request.form.get('chills')),
            'alcoholism': bool(request.form.get('alcoholism')),
            'hypertension': bool(request.form.get('hypertension')),
            'fainting': bool(request.form.get('fainting')),
            'dizziness': bool(request.form.get('dizziness')),
            'smoking': bool(request.form.get('smoking')),
            'obesity': bool(request.form.get('obesity')),
            'murmur': bool(request.form.get('murmur'))
        }
        
        # Prepare data for prediction
        prediction_data = {
            'Age': patient_data['age'],
            'Chest pain': int(patient_data['chest_pain']),
            'Shortness of breath': int(patient_data['shortness_of_breath']),
            'Fatigue': int(patient_data['fatigue']),
            'Systolic': patient_data['systolic'],
            'Diastolic': patient_data['diastolic'],
            'Heart rate (bpm)': patient_data['heart_rate'],
            'Lung sounds': int(patient_data['lung_sounds']),
            'Cholesterol level (mg/dL)': patient_data['cholesterol'],
            'LDL level (mg/dL)': patient_data['ldl'],
            'HDL level (mg/dL)': patient_data['hdl'],
            'Diabetes': int(patient_data['diabetes']),
            'Atrial fibrillation': int(patient_data['atrial_fibrillation']),
            'Rheumatic fever': int(patient_data['rheumatic_fever']),
            'Mitral stenosis': int(patient_data['mitral_stenosis']),
            'Aortic stenosis': int(patient_data['aortic_stenosis']),
            'Tricuspid stenosis': int(patient_data['tricuspid_stenosis']),
            'Pulmonary stenosis': int(patient_data['pulmonary_stenosis']),
            'Dilated cardiomyopathy': int(patient_data['dilated_cardiomyopathy']),
            'Hypertrophic cardiomyopathy': int(patient_data['hypertrophic_cardiomyopathy']),
            'Drug use': int(patient_data['drug_use']),
            'Fever': int(patient_data['fever']),
            'Chills': int(patient_data['chills']),
            'Alcoholism': int(patient_data['alcoholism']),
            'Hypertension': int(patient_data['hypertension']),
            'Fainting': int(patient_data['fainting']),
            'Dizziness': int(patient_data['dizziness']),
            'Smoking': int(patient_data['smoking']),
            'Obesity': int(patient_data['obesity']),
            'Murmur': int(patient_data['murmur'])
        }
        
        try:
            # Make prediction
            prediction_df = pd.DataFrame([prediction_data])
            prediction_result = model.predict(prediction_df)[0]
            
            # Create patient record
            patient = Patient(
                name=patient_data['name'],
                gender=patient_data['gender'],
                age=patient_data['age'],
                chest_pain=patient_data['chest_pain'],
                shortness_of_breath=patient_data['shortness_of_breath'],
                fatigue=patient_data['fatigue'],
                systolic=patient_data['systolic'],
                diastolic=patient_data['diastolic'],
                heart_rate=patient_data['heart_rate'],
                lung_sounds=patient_data['lung_sounds'],
                cholesterol=patient_data['cholesterol'],
                ldl=patient_data['ldl'],
                hdl=patient_data['hdl'],
                diabetes=patient_data['diabetes'],
                atrial_fibrillation=patient_data['atrial_fibrillation'],
                rheumatic_fever=patient_data['rheumatic_fever'],
                mitral_stenosis=patient_data['mitral_stenosis'],
                aortic_stenosis=patient_data['aortic_stenosis'],
                tricuspid_stenosis=patient_data['tricuspid_stenosis'],
                pulmonary_stenosis=patient_data['pulmonary_stenosis'],
                dilated_cardiomyopathy=patient_data['dilated_cardiomyopathy'],
                hypertrophic_cardiomyopathy=patient_data['hypertrophic_cardiomyopathy'],
                drug_use=patient_data['drug_use'],
                fever=patient_data['fever'],
                chills=patient_data['chills'],
                alcoholism=patient_data['alcoholism'],
                hypertension=patient_data['hypertension'],
                fainting=patient_data['fainting'],
                dizziness=patient_data['dizziness'],
                smoking=patient_data['smoking'],
                obesity=patient_data['obesity'],
                murmur=patient_data['murmur'],
                prediction_result=prediction_result,
                doctor_id=current_user.id
            )
            
            db.session.add(patient)
            db.session.commit()
            
            flash(f'Patient added successfully! Prediction: {prediction_result}', 'success')
            return redirect(url_for('doctor_dashboard'))
        except Exception as e:
            flash(f'Error making prediction: {str(e)}', 'error')
            return redirect(url_for('add_patient'))
    
    return render_template('doctor/add_patient.html')

# Dean Routes
@app.route('/dean/dashboard')
@login_required
def dean_dashboard():
    if current_user.role != 'dean':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.order_by(Patient.timestamp.desc()).all()
    doctors = User.query.filter_by(role='doctor').all()
    
    return render_template('dean/dashboard.html', patients=patients, doctors=doctors)

@app.route('/dean/register_doctor', methods=['GET', 'POST'])
@login_required
def register_doctor():
    if current_user.role != 'dean':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        username = request.form.get('username')
        email = request.form.get('email')
        password = request.form.get('password')
        
        if User.query.filter_by(username=username).first():
            flash('Username already exists', 'error')
        elif User.query.filter_by(email=email).first():
            flash('Email already exists', 'error')
        else:
            user = User(
                username=username,
                email=email,
                password_hash=generate_password_hash(password),
                role='doctor'
            )
            db.session.add(user)
            db.session.commit()
            flash('Doctor registered successfully!', 'success')
            return redirect(url_for('dean_dashboard'))
    
    return render_template('dean/register_doctor.html')

@app.route('/dean/analytics')
@login_required
def dean_analytics():
    if current_user.role != 'dean':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.all()
    
    # Analytics data
    total_patients = len(patients)
    predictions = [p.prediction_result for p in patients]
    prediction_counts = Counter(predictions)
    
    # Doctor-wise patient counts
    doctor_counts = {}
    for patient in patients:
        doctor_name = patient.doctor.username
        doctor_counts[doctor_name] = doctor_counts.get(doctor_name, 0) + 1
    
    return render_template('dean/analytics.html', 
                         total_patients=total_patients,
                         prediction_counts=prediction_counts,
                         doctor_counts=doctor_counts)

# Chairman Routes
@app.route('/chairman/dashboard')
@login_required
def chairman_dashboard():
    if current_user.role != 'chairman':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.order_by(Patient.timestamp.desc()).all()
    users = User.query.all()
    
    return render_template('chairman/dashboard.html', patients=patients, users=users)

@app.route('/chairman/register_user', methods=['GET', 'POST'])
@login_required
def register_user():
    if current_user.role != 'chairman':
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        username = request.form.get('username')
        email = request.form.get('email')
        password = request.form.get('password')
        role = request.form.get('role')
        
        if User.query.filter_by(username=username).first():
            flash('Username already exists', 'error')
        elif User.query.filter_by(email=email).first():
            flash('Email already exists', 'error')
        else:
            user = User(
                username=username,
                email=email,
                password_hash=generate_password_hash(password),
                role=role
            )
            db.session.add(user)
            db.session.commit()
            flash(f'{role.title()} registered successfully!', 'success')
            return redirect(url_for('chairman_dashboard'))
    
    return render_template('chairman/register_user.html')

# Export Routes
@app.route('/export/excel')
@login_required
def export_excel():
    if current_user.role not in ['dean', 'chairman']:
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Patient Records"
    
    # Headers
    headers = ['ID', 'Name', 'Gender', 'Age', 'Prediction', 'Doctor', 'Date']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Data
    for row, patient in enumerate(patients, 2):
        ws.cell(row=row, column=1, value=patient.id)
        ws.cell(row=row, column=2, value=patient.name)
        ws.cell(row=row, column=3, value=patient.gender)
        ws.cell(row=row, column=4, value=patient.age)
        ws.cell(row=row, column=5, value=patient.prediction_result)
        ws.cell(row=row, column=6, value=patient.doctor.username)
        ws.cell(row=row, column=7, value=patient.timestamp.strftime('%Y-%m-%d %H:%M'))
    
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='patient_records.xlsx'
    )

@app.route('/export/pdf')
@login_required
def export_pdf():
    if current_user.role not in ['dean', 'chairman']:
        flash('Access denied', 'error')
        return redirect(url_for('index'))
    
    patients = Patient.query.all()
    
    # Create PDF
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    
    # Table data
    data = [['ID', 'Name', 'Gender', 'Age', 'Prediction', 'Doctor', 'Date']]
    for patient in patients:
        data.append([
            str(patient.id),
            patient.name,
            patient.gender,
            str(patient.age),
            patient.prediction_result,
            patient.doctor.username,
            patient.timestamp.strftime('%Y-%m-%d %H:%M')
        ])
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    elements.append(table)
    doc.build(elements)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/pdf',
        as_attachment=True,
        download_name='patient_records.pdf'
    )

# API Routes for Analytics
@app.route('/api/analytics')
@login_required
def api_analytics():
    if current_user.role not in ['dean', 'chairman']:
        return jsonify({'error': 'Access denied'}), 403
    
    patients = Patient.query.all()
    
    # Prediction distribution
    predictions = [p.prediction_result for p in patients]
    prediction_counts = Counter(predictions)
    
    # Doctor-wise counts
    doctor_counts = {}
    for patient in patients:
        doctor_name = patient.doctor.username
        doctor_counts[doctor_name] = doctor_counts.get(doctor_name, 0) + 1
    
    return jsonify({
        'total_patients': len(patients),
        'prediction_counts': dict(prediction_counts),
        'doctor_counts': doctor_counts
    })

def init_database():
    """Initialize database with retry logic"""
    max_retries = 3
    retry_delay = 2
    
    for attempt in range(max_retries):
        try:
            with app.app_context():
                db.create_all()
                
                # Create default chairman if none exists
                if not User.query.filter_by(role='chairman').first():
                    chairman = User(
                        username='admin',
                        email='admin@hospital.com',
                        password_hash=generate_password_hash('admin123'),
                        role='chairman'
                    )
                    db.session.add(chairman)
                    db.session.commit()
                    print("Default chairman created: username='admin', password='admin123'")
                else:
                    print("Database already initialized")
                return True
        except Exception as e:
            print(f"Database initialization attempt {attempt + 1} failed: {e}")
            if attempt < max_retries - 1:
                print(f"Retrying in {retry_delay} seconds...")
                time.sleep(retry_delay)
                retry_delay *= 2
            else:
                print("Failed to initialize database after all retries")
                return False

if __name__ == '__main__':
    # Initialize database
    if init_database():
        app.run(debug=True)
    else:
        print("Please restart MySQL service in XAMPP and try again")
