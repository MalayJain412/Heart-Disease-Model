from flask import Blueprint, send_file, flash, redirect, url_for
from flask_login import login_required, current_user
from models import Patient
import io
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors

export_bp = Blueprint('export', __name__)

@export_bp.route('/export/excel')
@login_required
def export_excel():
    if current_user.role not in ['dean', 'chairman']:
        flash('Access denied', 'error')
        return redirect(url_for('main.index'))
    patients = Patient.query.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Patient Records"
    headers = ['ID', 'Name', 'Gender', 'Age', 'Prediction', 'Doctor', 'Date']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    for row, patient in enumerate(patients, 2):
        ws.cell(row=row, column=1, value=patient.id)
        ws.cell(row=row, column=2, value=patient.name)
        ws.cell(row=row, column=3, value=patient.gender)
        ws.cell(row=row, column=4, value=patient.age)
        ws.cell(row=row, column=5, value=patient.prediction_result)
        ws.cell(row=row, column=6, value=patient.doctor.username)
        ws.cell(row=row, column=7, value=patient.timestamp.strftime('%Y-%m-%d %H:%M'))
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='patient_records.xlsx'
    )

@export_bp.route('/export/pdf')
@login_required
def export_pdf():
    if current_user.role not in ['dean', 'chairman']:
        flash('Access denied', 'error')
        return redirect(url_for('main.index'))
    patients = Patient.query.all()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    data = [['ID', 'Name', 'Gender', 'Age', 'Prediction', 'Doctor', 'Date']]
    for patient in patients:
        data.append([
            str(patient.id),
            patient.name,
            patient.gender,
            str(patient.age),
            patient.prediction_result,
            patient.doctor.username,
            patient.timestamp.strftime('%Y-%m-%d %H:%M')
        ])
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(table)
    doc.build(elements)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/pdf',
        as_attachment=True,
        download_name='patient_records.pdf'
    ) 